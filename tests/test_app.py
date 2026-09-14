import unittest
import json
import os
import sys
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app, db
from app.models.user import User
from app.models.resume import Resume
from app.models.job import Job
from app.models.match import Match


class BaseTestCase(unittest.TestCase):
    def setUp(self):
        self.app = create_app('testing')
        self.client = self.app.test_client()
        with self.app.app_context():
            db.create_all()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()

    def register_user(self, name='Test User', email='test@test.com', password='password123', role='job_seeker'):
        return self.client.post('/register', data={
            'name': name, 'email': email, 'password': password,
            'confirm_password': password, 'role': role
        }, follow_redirects=False)

    def login_user(self, email='test@test.com', password='password123'):
        return self.client.post('/login', data={
            'email': email, 'password': password
        }, follow_redirects=False)


class AuthTestCase(BaseTestCase):
    def test_registration(self):
        resp = self.register_user()
        self.assertEqual(resp.status_code, 302)
        with self.app.app_context():
            user = User.query.filter_by(email='test@test.com').first()
            self.assertIsNotNone(user)
            self.assertEqual(user.role, 'job_seeker')

    def test_duplicate_email(self):
        self.register_user()
        resp = self.register_user()
        self.assertEqual(resp.status_code, 200)

    def test_login(self):
        self.register_user()
        resp = self.login_user()
        self.assertEqual(resp.status_code, 302)

    def test_invalid_login(self):
        resp = self.login_user(password='wrong')
        self.assertEqual(resp.status_code, 200)

    def test_logout(self):
        self.register_user()
        self.login_user()
        resp = self.client.get('/logout', follow_redirects=True)
        self.assertEqual(resp.status_code, 200)

    def test_role_based_access(self):
        self.register_user(name='Seeker', email='s@t.com', role='job_seeker')
        self.login_user(email='s@t.com')
        resp = self.client.get('/recruiter/dashboard')
        self.assertEqual(resp.status_code, 302)

        self.client.get('/logout')
        self.register_user(name='Rec', email='r@t.com', role='recruiter')
        self.login_user(email='r@t.com')
        resp = self.client.get('/seeker/dashboard')
        self.assertEqual(resp.status_code, 302)


class ResumeTestCase(BaseTestCase):
    def setUp(self):
        super().setUp()
        self.register_user(name='Seeker', email='s@t.com', role='job_seeker')
        self.login_user(email='s@t.com')

    def test_upload_docx(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('John Smith')
        doc.add_paragraph('Email: john@test.com')
        doc.add_paragraph('Skills: Python, Flask, SQL')
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        resp = self.client.post('/resume/upload', data={
            'resume': (buf, 'resume.docx')
        }, content_type='multipart/form-data', follow_redirects=True)
        self.assertEqual(resp.status_code, 200)
        with self.app.app_context():
            resume = Resume.query.first()
            self.assertIsNotNone(resume)
            self.assertEqual(resume.candidate_name, 'John Smith')

    def test_upload_invalid_file(self):
        resp = self.client.post('/resume/upload', data={
            'resume': (BytesIO(b'test'), 'test.txt')
        }, content_type='multipart/form-data', follow_redirects=True)
        self.assertEqual(resp.status_code, 200)


class NLPTestCase(BaseTestCase):
    def test_email_extraction(self):
        from app.services.nlp_service import extract_email
        self.assertEqual(extract_email('Contact: test@example.com'), 'test@example.com')
        self.assertEqual(extract_email('No email here'), '')

    def test_phone_extraction(self):
        from app.services.nlp_service import extract_phone
        phone = extract_phone('Call +234 801 234 5678')
        self.assertIn('234', phone)

    def test_skills_extraction(self):
        from app.services.nlp_service import extract_skills
        skills = extract_skills('Python Flask SQL Git Docker')
        self.assertIn('python', skills)
        self.assertIn('flask', skills)
        self.assertIn('sql', skills)

    def test_education_extraction(self):
        from app.services.nlp_service import extract_education
        edu = extract_education('B.Sc. Computer Science\nUniversity of Lagos\n2020')
        self.assertTrue(len(edu) > 0)
        self.assertIn('B.Sc', edu[0]['degree'])


class LegacyMatchingTestCase(BaseTestCase):
    """Tests for the legacy matching_service.py functions (kept for backward compatibility)."""

    def test_tfidf_similarity(self):
        from app.services.matching_service import calculate_tfidf_similarity
        score = calculate_tfidf_similarity(
            'Python Flask developer',
            'Looking for a Python developer with Flask'
        )
        self.assertGreater(score, 0)

    def test_skill_matching(self):
        from app.services.matching_service import calculate_skill_score
        score, matched, missing = calculate_skill_score(
            ['Python', 'Flask', 'SQL'],
            ['Python', 'Flask', 'SQL', 'Git', 'Docker']
        )
        self.assertEqual(score, 60.0)
        self.assertEqual(len(matched), 3)
        self.assertEqual(len(missing), 2)

    def test_experience_score(self):
        from app.services.matching_service import calculate_experience_score
        self.assertEqual(calculate_experience_score(5, 3), 100.0)
        self.assertEqual(calculate_experience_score(2, 3), 66.67)

    def test_final_score(self):
        from app.services.matching_service import calculate_final_score
        score = calculate_final_score(80, 80, 100, 100)
        self.assertGreater(score, 0)
        self.assertLessEqual(score, 100)


class EmbeddingServiceTestCase(BaseTestCase):
    """Tests for content hashing and embedding generation."""

    def test_content_hash_deterministic(self):
        from app.services.embedding_service import compute_content_hash
        h1 = compute_content_hash('Hello World')
        h2 = compute_content_hash('Hello World')
        self.assertEqual(h1, h2)

    def test_content_hash_different_for_different_content(self):
        from app.services.embedding_service import compute_content_hash
        h1 = compute_content_hash('Hello World')
        h2 = compute_content_hash('Goodbye World')
        self.assertNotEqual(h1, h2)

    def test_content_hash_none_for_empty(self):
        from app.services.embedding_service import compute_content_hash
        self.assertIsNone(compute_content_hash(''))
        self.assertIsNone(compute_content_hash(None))

    def test_serialize_deserialize_embedding(self):
        from app.services.embedding_service import serialize_embedding, deserialize_embedding
        import numpy as np
        original = np.array([0.5, 0.3, 0.8], dtype=np.float32)
        serialized = serialize_embedding(original)
        self.assertIsNotNone(serialized)
        deserialized = deserialize_embedding(serialized)
        self.assertIsNotNone(deserialized)
        np.testing.assert_array_almost_equal(original, deserialized, decimal=5)

    def test_generate_embedding_from_analysis(self):
        from app.services.embedding_service import generate_embedding_from_analysis
        analysis = {
            'primary_roles': ['Python Developer'],
            'direct_skills': ['Python', 'Flask'],
            'inferred_capabilities': ['API design'],
        }
        embedding = generate_embedding_from_analysis(analysis)
        self.assertIsNotNone(embedding)
        self.assertGreater(len(embedding), 0)

    def test_generate_embedding_none_for_empty(self):
        from app.services.embedding_service import generate_embedding_from_analysis
        self.assertIsNone(generate_embedding_from_analysis(None))
        self.assertIsNone(generate_embedding_from_analysis({}))

    def test_cosine_similarity(self):
        from app.services.embedding_service import cosine_similarity_vec
        import numpy as np
        a = np.array([1.0, 0.0, 0.0])
        b = np.array([1.0, 0.0, 0.0])
        self.assertAlmostEqual(cosine_similarity_vec(a, b), 1.0)

        c = np.array([0.0, 1.0, 0.0])
        self.assertAlmostEqual(cosine_similarity_vec(a, c), 0.0)

    def test_cosine_similarity_none(self):
        from app.services.embedding_service import cosine_similarity_vec
        self.assertEqual(cosine_similarity_vec(None, None), 0.0)


class SemanticMatchingTestCase(BaseTestCase):
    """Tests for the semantic matching service."""

    def test_skill_matches_direct(self):
        from app.services.semantic_matching_service import compute_skill_matches
        job_analysis = {
            'critical_skills': ['Python', 'Flask'],
            'important_skills': ['SQL'],
            'preferred_skills': ['Docker'],
        }
        result = compute_skill_matches(['python', 'flask', 'sql', 'git'], job_analysis)
        self.assertIn('python', result['direct'])
        self.assertIn('flask', result['direct'])
        self.assertIn('sql', result['direct'])
        self.assertEqual(result['total_required'], 3)

    def test_skill_matches_related(self):
        """Django developer should match Python requirement via relationships."""
        from app.services.semantic_matching_service import compute_skill_matches
        job_analysis = {
            'critical_skills': ['Python'],
            'important_skills': [],
            'preferred_skills': [],
        }
        result = compute_skill_matches(['django', 'html'], job_analysis)
        # Django is related to Python via SKILL_RELATIONSHIPS
        self.assertTrue(
            len(result['related']) > 0 or len(result['direct']) > 0,
            "Django should match Python as a related skill"
        )

    def test_skill_matches_missing(self):
        from app.services.semantic_matching_service import compute_skill_matches
        job_analysis = {
            'critical_skills': ['Python', 'Docker', 'Kubernetes'],
            'important_skills': [],
            'preferred_skills': [],
        }
        result = compute_skill_matches(['python'], job_analysis)
        self.assertIn('python', result['direct'])
        self.assertIn('docker', result['missing'])
        self.assertIn('kubernetes', result['missing'])

    def test_skill_score_perfect(self):
        from app.services.semantic_matching_service import compute_skill_score
        matches = {'direct': ['a', 'b'], 'related': [], 'missing': [], 'total_required': 2}
        self.assertEqual(compute_skill_score(matches), 100.0)

    def test_skill_score_with_related(self):
        from app.services.semantic_matching_service import compute_skill_score
        # 1 direct + 1 related out of 2 required = (1*1.0 + 1*0.6) / 2 = 80%
        matches = {'direct': ['a'], 'related': ['b'], 'missing': [], 'total_required': 2}
        self.assertEqual(compute_skill_score(matches), 80.0)

    def test_skill_score_no_requirements(self):
        from app.services.semantic_matching_service import compute_skill_score
        matches = {'direct': [], 'related': [], 'missing': [], 'total_required': 0}
        self.assertEqual(compute_skill_score(matches), 100.0)

    def test_capability_score(self):
        from app.services.semantic_matching_service import compute_capability_score
        candidate = {'inferred_capabilities': ['API design', 'team leadership']}
        job = {'inferred_capabilities': ['API design', 'database management', 'team leadership']}
        score = compute_capability_score(candidate, job)
        self.assertGreater(score, 50.0)

    def test_experience_score_exact_match(self):
        from app.services.semantic_matching_service import compute_experience_score
        job = {'experience_level': 'mid', 'min_years_experience': 3}
        score = compute_experience_score(3, 'mid', job)
        self.assertEqual(score, 100.0)

    def test_experience_score_underqualified(self):
        from app.services.semantic_matching_service import compute_experience_score
        job = {'experience_level': 'senior', 'min_years_experience': 5}
        score = compute_experience_score(2, 'entry', job)
        self.assertLess(score, 60.0)

    def test_education_score_meets_requirement(self):
        from app.services.semantic_matching_service import compute_education_score
        job = {'education_requirement': 'bachelors'}
        score = compute_education_score('masters', job)
        self.assertEqual(score, 100.0)

    def test_education_score_no_requirement(self):
        from app.services.semantic_matching_service import compute_education_score
        job = {'education_requirement': ''}
        score = compute_education_score('bachelors', job)
        self.assertGreater(score, 80.0)

    def test_final_score_weighted(self):
        from app.services.semantic_matching_service import compute_final_score
        score = compute_final_score(
            semantic_score=80.0,
            skill_score=90.0,
            capability_score=70.0,
            experience_score=100.0,
            education_score=100.0
        )
        # weighted: 80*0.45 + 90*0.25 + 70*0.15 + 100*0.10 + 100*0.05
        # = 36 + 22.5 + 10.5 + 10 + 5 = 84.0
        self.assertAlmostEqual(score, 84.0, delta=0.2)

    def test_skill_relationships_cover_key_mappings(self):
        from app.services.semantic_matching_service import SKILL_RELATIONSHIPS
        # Django should be related to Python
        self.assertIn('django', SKILL_RELATIONSHIPS)
        self.assertIn('python', SKILL_RELATIONSHIPS['django'])
        # React should be related to JavaScript
        self.assertIn('react', SKILL_RELATIONSHIPS)
        self.assertIn('javascript', SKILL_RELATIONSHIPS['react'])

    def test_django_matches_python_job(self):
        """Semantic test: Django Developer applying to Python job should score high."""
        from app.services.semantic_matching_service import compute_skill_matches, compute_skill_score
        job_analysis = {
            'critical_skills': ['python', 'flask', 'sql'],
            'important_skills': ['git'],
            'preferred_skills': ['docker'],
        }
        # Django developer's skills
        candidate_skills = ['django', 'python', 'html', 'css', 'sql', 'git']
        matches = compute_skill_matches(candidate_skills, job_analysis)
        score = compute_skill_score(matches)
        # Django + Python both match Python (direct), SQL matches, Git matches
        # Should be well above 50%
        self.assertGreater(score, 50.0)

    def test_react_matches_js_frontend_job(self):
        """Semantic test: React developer applying to JS Frontend job should score high."""
        from app.services.semantic_matching_service import compute_skill_matches, compute_skill_score
        job_analysis = {
            'critical_skills': ['javascript', 'html', 'css'],
            'important_skills': ['react'],
            'preferred_skills': ['vue.js'],
        }
        candidate_skills = ['react', 'javascript', 'html', 'css', 'git']
        matches = compute_skill_matches(candidate_skills, job_analysis)
        score = compute_skill_score(matches)
        self.assertGreater(score, 60.0)

    def test_property_manager_low_for_python_dev(self):
        """Semantic test: Property Manager applying to Python Developer should score low."""
        from app.services.semantic_matching_service import compute_skill_matches, compute_skill_score
        job_analysis = {
            'critical_skills': ['python', 'flask', 'sql'],
            'important_skills': ['git', 'docker'],
            'preferred_skills': ['aws'],
        }
        # Property Manager's skills
        candidate_skills = ['property management', 'tenant relations', 'lease administration', ' microsoft office']
        matches = compute_skill_matches(candidate_skills, job_analysis)
        score = compute_skill_score(matches)
        # Should be very low - no overlap
        self.assertLess(score, 20.0)


class JobTestCase(BaseTestCase):
    def setUp(self):
        super().setUp()
        self.register_user(name='Recruiter', email='r@t.com', role='recruiter')
        self.login_user(email='r@t.com')

    def test_create_job(self):
        resp = self.client.post('/recruiter/jobs/create', data={
            'title': 'Python Developer', 'company': 'TestCo',
            'description': 'A great job', 'required_skills': 'Python, Flask',
            'minimum_experience': '2', 'employment_type': 'full-time'
        }, follow_redirects=False)
        self.assertEqual(resp.status_code, 302)
        with self.app.app_context():
            job = Job.query.first()
            self.assertIsNotNone(job)
            self.assertEqual(job.title, 'Python Developer')
            # New fields should have defaults
            self.assertEqual(job.analysis_status, 'pending')

    def test_edit_job(self):
        self.client.post('/recruiter/jobs/create', data={
            'title': 'Original', 'company': 'TestCo',
            'description': 'Desc', 'employment_type': 'full-time'
        })
        with self.app.app_context():
            job = Job.query.first()
            resp = self.client.post(f'/recruiter/jobs/{job.id}/edit', data={
                'title': 'Updated', 'company': 'TestCo',
                'description': 'Desc', 'employment_type': 'full-time'
            }, follow_redirects=False)
            self.assertEqual(resp.status_code, 302)
            job = Job.query.get(job.id)
            self.assertEqual(job.title, 'Updated')

    def test_close_job(self):
        self.client.post('/recruiter/jobs/create', data={
            'title': 'Job', 'company': 'TestCo',
            'description': 'Desc', 'employment_type': 'full-time'
        })
        with self.app.app_context():
            job = Job.query.first()
            resp = self.client.post(f'/recruiter/jobs/{job.id}/close', follow_redirects=False)
            self.assertEqual(resp.status_code, 302)
            job = Job.query.get(job.id)
            self.assertEqual(job.status, 'closed')

    def test_delete_job(self):
        self.client.post('/recruiter/jobs/create', data={
            'title': 'Job', 'company': 'TestCo',
            'description': 'Desc', 'employment_type': 'full-time'
        })
        with self.app.app_context():
            job = Job.query.first()
            resp = self.client.post(f'/recruiter/jobs/{job.id}/delete', follow_redirects=False)
            self.assertEqual(resp.status_code, 302)
            self.assertIsNone(Job.query.get(job.id))


class ModelFieldTestCase(BaseTestCase):
    """Test that new model fields exist and have correct defaults."""

    def test_resume_new_fields(self):
        with self.app.app_context():
            user = User(name='Test', email='t@t.com', role='job_seeker')
            user.set_password('pass')
            db.session.add(user)
            db.session.commit()

            resume = Resume(
                user_id=user.id, filename='test.docx', file_type='docx',
                raw_text='test', analysis_status='completed',
                content_hash='abc123', analysis_version=1
            )
            db.session.add(resume)
            db.session.commit()

            loaded = Resume.query.get(resume.id)
            self.assertEqual(loaded.analysis_status, 'completed')
            self.assertEqual(loaded.content_hash, 'abc123')
            self.assertEqual(loaded.analysis_version, 1)
            self.assertIsNone(loaded.semantic_analysis)
            self.assertIsNone(loaded.semantic_embedding)

    def test_job_new_fields(self):
        with self.app.app_context():
            user = User(name='Rec', email='r@r.com', role='recruiter')
            user.set_password('pass')
            db.session.add(user)
            db.session.commit()

            job = Job(
                recruiter_id=user.id, title='Test', company='Co',
                description='Desc', analysis_status='completed',
                content_hash='def456'
            )
            db.session.add(job)
            db.session.commit()

            loaded = Job.query.get(job.id)
            self.assertEqual(loaded.analysis_status, 'completed')
            self.assertEqual(loaded.content_hash, 'def456')

    def test_match_new_fields(self):
        with self.app.app_context():
            u = User(name='U', email='u@u.com', role='job_seeker')
            u.set_password('p')
            db.session.add(u)
            db.session.commit()

            r = Resume(user_id=u.id, filename='r.docx', file_type='docx')
            j = Job(recruiter_id=u.id, title='J', company='C', description='D')
            db.session.add_all([r, j])
            db.session.commit()

            m = Match(
                resume_id=r.id, job_id=j.id,
                semantic_score=75.0, skill_score=80.0, capability_score=70.0,
                experience_score=90.0, education_score=100.0, final_score=82.0,
                direct_matches=json.dumps(['python']),
                related_matches=json.dumps(['django']),
                missing_skills=json.dumps(['docker']),
                explanation='Good match.',
                analysis_version='r1_j1'
            )
            db.session.add(m)
            db.session.commit()

            loaded = Match.query.get(m.id)
            self.assertEqual(loaded.semantic_score, 75.0)
            self.assertEqual(loaded.capability_score, 70.0)
            self.assertEqual(loaded.analysis_version, 'r1_j1')
            self.assertEqual(json.loads(loaded.direct_matches), ['python'])


if __name__ == '__main__':
    unittest.main()
